from __future__ import annotations

from io import BytesIO
from pathlib import Path
import html
import logging
import os
import tempfile

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Mm, Pt, RGBColor

from .models import Resume
from .document_converter import ConversionError, convert_document


NAVY = RGBColor(24, 43, 74)
SLATE = RGBColor(91, 106, 126)
ACCENT = RGBColor(53, 115, 219)


def _template_font() -> str:
    return os.environ.get("PAPERPRO_RESUME_FONT", "").strip() or (
        "Microsoft YaHei" if os.name == "nt" else "Noto Sans CJK SC"
    )


def _set_font(run, *, size: float, bold: bool = False, color=NAVY) -> None:
    font = _template_font()
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    fonts = run._element.get_or_add_rPr().rFonts
    if fonts is not None:
        fonts.set(qn("w:eastAsia"), font)


def _heading(doc: Document, name: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(name)
    _set_font(run, size=11, bold=True)
    border = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "7")
    bottom.set(qn("w:color"), "D5DFEE")
    border.append(bottom)
    p._p.get_or_add_pPr().append(border)


def _paragraph(doc: Document, content: str, *, bold: bool = False, bullet: bool = False) -> None:
    p = doc.add_paragraph(style="List Bullet" if bullet else None)
    p.paragraph_format.space_after = Pt(2.2)
    p.paragraph_format.line_spacing = 1.13
    if bullet:
        p.paragraph_format.left_indent = Inches(.22)
        p.paragraph_format.first_line_indent = Inches(-.1)
    run = p.add_run(content)
    _set_font(run, size=9.1, bold=bold, color=NAVY if bold else SLATE)


def make_docx(resume: Resume) -> bytes:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Inches(.68)
    section.bottom_margin = Inches(.58)
    section.left_margin = Inches(.73)
    section.right_margin = Inches(.73)
    section.header_distance = Inches(.32)
    section.footer_distance = Inches(.32)
    normal = doc.styles["Normal"]
    normal.font.name = _template_font()
    normal.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), _template_font())
    normal.font.size = Pt(9.1)
    normal.paragraph_format.space_after = Pt(0)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    _set_font(p.add_run(resume.name.strip() or "未填写姓名"), size=21, bold=True)
    if resume.title.strip():
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        _set_font(p.add_run(resume.title.strip()), size=10, color=ACCENT, bold=True)
    contact = "   |   ".join(s for s in [resume.phone, resume.email, resume.location, *resume.links] if s.strip())
    if contact:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        _set_font(p.add_run(contact), size=8.3, color=SLATE)

    if resume.summary.strip():
        _heading(doc, "个人简介  /  PROFILE")
        for line in resume.summary.splitlines():
            if line.strip():
                _paragraph(doc, line.strip())
    if resume.skills:
        _heading(doc, "专业技能  /  SKILLS")
        _paragraph(doc, "  ·  ".join(s.strip() for s in resume.skills if s.strip()))
    if resume.experience:
        _heading(doc, "工作经历  /  EXPERIENCE")
        for exp in resume.experience:
            if exp.heading.strip():
                _paragraph(doc, exp.heading.strip(), bold=True)
            for bullet in exp.bullets:
                if bullet.strip():
                    _paragraph(doc, bullet.strip(), bullet=True)
    if resume.projects:
        _heading(doc, "项目经历  /  PROJECTS")
        for project in resume.projects:
            if project.heading.strip():
                _paragraph(doc, project.heading.strip(), bold=True)
            for bullet in project.bullets:
                if bullet.strip():
                    _paragraph(doc, bullet.strip(), bullet=True)
    if resume.education:
        _heading(doc, "教育背景  /  EDUCATION")
        for entry in resume.education:
            if entry.heading.strip():
                _paragraph(doc, entry.heading.strip(), bold=True)
            for detail in entry.details:
                if detail.strip():
                    _paragraph(doc, detail.strip())
    if resume.additional:
        _heading(doc, "其他信息  /  ADDITIONAL")
        for item in resume.additional:
            if item.strip():
                _paragraph(doc, item.strip(), bullet=True)

    doc.core_properties.title = "Resume"
    doc.core_properties.author = ""
    doc.core_properties.subject = ""
    doc.core_properties.keywords = ""
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _word_pdf(docx_data: bytes) -> bytes:
    import pythoncom
    import win32com.client
    with tempfile.TemporaryDirectory(prefix="paperpro-export-") as temp:
        source = Path(temp) / "resume.docx"
        dest = Path(temp) / "resume.pdf"
        source.write_bytes(docx_data)
        pythoncom.CoInitialize()
        word = None
        document = None
        try:
            word = win32com.client.DispatchEx("Word.Application")
            word.Visible = False
            word.DisplayAlerts = 0
            word.AutomationSecurity = 3
            document = word.Documents.Open(str(source), ReadOnly=True, AddToRecentFiles=False, ConfirmConversions=False)
            document.ExportAsFixedFormat(str(dest), 17)
            if not dest.exists() or dest.stat().st_size == 0:
                raise RuntimeError("Word 未创建 PDF")
            return dest.read_bytes()
        finally:
            try:
                if document is not None:
                    try:
                        document.Close(SaveChanges=False)
                    except Exception as error:
                        logging.warning("Word document cleanup failed: %s", type(error).__name__)
                if word is not None:
                    try:
                        word.Quit()
                    except Exception as error:
                        logging.warning("Word application cleanup failed: %s", type(error).__name__)
            finally:
                pythoncom.CoUninitialize()


def _fallback_pdf(resume: Resume) -> bytes:
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.lib import colors
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, HRFlowable
    # ReportLab's TrueType fallback cannot encode non-BMP glyphs correctly.
    # Never return a PDF with a corrupted candidate name: Word handles this path.
    if any(ord(char) > 0xFFFF for char in str(resume.model_dump())):
        raise ValueError("当前备用 PDF 引擎无法保证罕见汉字不丢失，请使用 Word/LibreOffice 导出或检查转换服务")
    main_font = "STSong-Light"
    font_dir = Path("C:/Windows/Fonts")
    if (font_dir / "simhei.ttf").exists():
        pdfmetrics.registerFont(TTFont("PPChinese", str(font_dir / "simhei.ttf")))
        main_font = "PPChinese"
    elif Path("/usr/share/fonts/truetype/wqy/wqy-zenhei.ttc").exists():
        pdfmetrics.registerFont(TTFont("PPChinese", "/usr/share/fonts/truetype/wqy/wqy-zenhei.ttc", subfontIndex=0))
        main_font = "PPChinese"
    else:
        pdfmetrics.registerFont(UnicodeCIDFont(main_font))
    buff = BytesIO()
    pdf = SimpleDocTemplate(
        buff, pagesize=(595.2756, 841.8898), leftMargin=52, rightMargin=52,
        topMargin=45, bottomMargin=46,
        title="Resume", author=""
    )
    main = ParagraphStyle("main", fontName=main_font, fontSize=9, leading=14, wordWrap="CJK", textColor=colors.HexColor("#334155"))
    bold = ParagraphStyle("bold", parent=main, fontSize=10.3, leading=16, textColor=colors.HexColor("#182B4A"))
    title = ParagraphStyle("title", parent=bold, fontSize=22, leading=29)
    small = ParagraphStyle("small", parent=main, fontSize=8, leading=11)
    story = []

    def line(value: str, style=main):
        story.append(Paragraph(html.escape(value).replace("\n", "<br/>"), style))
        story.append(Spacer(1, 3))

    def section(label: str):
        story.append(Spacer(1, 8))
        line(label, bold)
        story.append(HRFlowable(width="100%", thickness=.6, color=colors.HexColor("#CED9E9")))
        story.append(Spacer(1, 5))

    line(resume.name or "未填写姓名", title)
    if resume.title:
        line(resume.title, bold)
    contact = " | ".join(v for v in [resume.phone, resume.email, resume.location, *resume.links] if v)
    if contact:
        line(contact, small)
    if resume.summary:
        section("个人简介 / PROFILE")
        line(resume.summary)
    if resume.skills:
        section("专业技能 / SKILLS")
        line("  ·  ".join(resume.skills))
    for label, entries in (("工作经历 / EXPERIENCE", resume.experience), ("项目经历 / PROJECTS", resume.projects)):
        if entries:
            section(label)
            for entry in entries:
                line(entry.heading, bold)
                for bullet in entry.bullets:
                    line("•  " + bullet)
    if resume.education:
        section("教育背景 / EDUCATION")
        for item in resume.education:
            line(item.heading, bold)
            for detail in item.details:
                line(detail)
    if resume.additional:
        section("其他信息 / ADDITIONAL")
        for item in resume.additional:
            line("•  " + item)
    pdf.build(story)
    return buff.getvalue()


def make_pdf(resume: Resume, docx_data: bytes | None = None) -> bytes:
    engine = os.environ.get("PAPERPRO_PDF_ENGINE", "auto")
    if engine not in {"auto", "libreoffice", "word", "reportlab"}:
        raise ValueError("PAPERPRO_PDF_ENGINE 必须是 auto、libreoffice、word 或 reportlab")
    if engine == "reportlab":
        return _fallback_pdf(resume)
    document = docx_data if docx_data is not None else make_docx(resume)
    if engine == "libreoffice":
        return convert_document(document, "docx", "pdf")
    if engine == "word" and os.name != "nt":
        raise ValueError("Word PDF 引擎仅支持 Windows，请使用 LibreOffice")
    if os.name == "nt":
        try:
            return _word_pdf(document)
        except Exception as exc:
            if engine == "word":
                raise
            logging.warning("Word PDF conversion failed; using safe PDF renderer: %s", type(exc).__name__)
    else:
        try:
            return convert_document(document, "docx", "pdf")
        except ConversionError as exc:
            logging.warning("LibreOffice PDF conversion failed; using safe PDF renderer: %s", type(exc).__name__)
    return _fallback_pdf(resume)

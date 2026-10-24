from io import BytesIO
from pathlib import Path

import fitz
import pytest
from docx import Document
from fastapi.testclient import TestClient
from PIL import Image, ImageDraw, ImageFont

from app.main import app
from app.models import Resume
from app.parser import parse_resume, extract_file
from app.enhancer import local_enhance, validate_ai_rewrite
from app.exporter import make_docx, make_pdf


client = TestClient(app)

SAMPLE = """陈晓明
chen@example.com  13812345678
求职意向：后端开发工程师
个人简介
拥有3年软件项目开发经验，熟悉Python。
专业技能
Python、FastAPI、MySQL
工作经历
2022.07-2025.06 星辰科技 | 后端开发工程师
- 开发订单管理接口，降低接口响应时间20%
- 负责数据库设计
教育背景
2018.09-2022.06 华南理工大学 | 计算机科学学士
项目经历
订单分析平台
- 使用Python构建分析模块
其他经历
校园志愿服务100小时
"""


def test_parse_chinese_preserves_facts_and_unknown():
    result = parse_resume(SAMPLE)
    assert result.name == "陈晓明"
    assert result.email == "chen@example.com"
    assert result.phone == "13812345678"
    assert "Python" in " ".join(result.skills)
    assert result.experience and "20%" in " ".join(result.experience[0].bullets)
    assert result.education and "华南理工大学" in result.education[0].heading
    assert "志愿服务100小时" in " ".join(result.additional)


def test_parse_english_and_table_like_fields():
    raw = """Alex Chen
alex@example.com · +1 (415) 555-1212
PROFESSIONAL SUMMARY
Builds web platforms.
SKILLS
Python, SQL, React
EXPERIENCE
Data Engineer | Example Labs | 2021 - 2024
• Improved daily jobs by 12%
EDUCATION
Example University | BSc | 2017 - 2021
"""
    result = parse_resume(raw)
    assert result.name == "Alex Chen"
    assert result.email == "alex@example.com"
    assert result.experience and "12%" in " ".join(result.experience[0].bullets)
    assert result.education and "Example University" in result.education[0].heading


def test_local_enhancement_preserves_numbers_and_identity():
    draft = parse_resume(SAMPLE)
    enhanced, changes = local_enhance(draft)
    assert enhanced.name == draft.name
    assert enhanced.email == draft.email
    assert "20%" in str(enhanced.model_dump())
    assert isinstance(changes, list)


def test_ai_rewrite_keeps_identity_and_flags_new_numeric_claims_for_human_review():
    from app.enhancer import _changes
    source = parse_resume(SAMPLE)
    proposal = {
        "summary": "拥有10年经验，曾获国际冠军",
        "experience_bullets": [["降低接口响应时间90%", "负责数据库设计"]],
        "project_bullets": [["使用Python构建分析模块"]],
        "name": "AI 编造的名字",
    }
    suggested = validate_ai_rewrite(source, proposal)
    assert suggested.name == source.name
    assert source.summary != suggested.summary
    flagged = _changes(source, suggested)
    assert any("数字" in warning for change in flagged for warning in change["review_flags"])
    assert source.model_dump() == parse_resume(SAMPLE).model_dump()


def test_docx_and_pdf_have_expected_text():
    resume = parse_resume(SAMPLE)
    docx = make_docx(resume)
    reopened = Document(BytesIO(docx))
    all_text = "\n".join(p.text for p in reopened.paragraphs)
    assert "陈晓明" in all_text
    assert "星辰科技" in all_text

    pdf = make_pdf(resume, docx)
    page_text = "\n".join(page.get_text() for page in fitz.open(stream=pdf, filetype="pdf"))
    assert "陈晓明" in page_text
    assert "星辰科技" in page_text


def test_extract_docx_table_content():
    doc = Document()
    doc.add_paragraph("Jordan Lee")
    doc.add_paragraph("jordan@example.com")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Skills"
    table.cell(0, 1).text = "Rust"
    blob = BytesIO()
    doc.save(blob)
    text, warnings, method = extract_file("candidate.docx", blob.getvalue())
    assert "Jordan Lee" in text
    assert "Rust" in text
    assert not warnings


def test_extract_scanned_png(tmp_path):
    img = Image.new("RGB", (1150, 250), "white")
    font = ImageFont.truetype("C:/Windows/Fonts/arial.ttf", 58)
    ImageDraw.Draw(img).text((30, 50), "ALEX CHEN RESUME", font=font, fill="black")
    stream = BytesIO()
    img.save(stream, "PNG")
    text, warnings, method = extract_file("scan.png", stream.getvalue())
    assert "ALEX" in text.upper()
    assert "ocr" in method.lower()


def test_extract_scanned_pdf():
    document = fitz.open()
    page = document.new_page(width=850, height=260)
    page.insert_text((65, 120), "CANDIDATE EXPERIENCE", fontsize=35)
    raster = page.get_pixmap(matrix=fitz.Matrix(1.8, 1.8))
    scanned = fitz.open()
    imgpage = scanned.new_page(width=850, height=260)
    imgpage.insert_image(imgpage.rect, stream=raster.tobytes("png"))
    text, warnings, method = extract_file("scan.pdf", scanned.tobytes())
    assert "CANDIDATE" in text.upper()
    assert "ocr" in method.lower()


def test_http_upload_rejects_bad_and_oversize():
    r = client.post("/api/extract", files={"file": ("bad.pdf", b"not pdf", "application/pdf")})
    assert r.status_code == 422
    r = client.post("/api/extract", files={"file": ("x.txt", b"test", "text/plain")})
    assert r.status_code == 415
    r = client.post("/api/extract", files={"file": ("big.png", b"x" * (12 * 1024 * 1024 + 1), "image/png")})
    assert r.status_code == 413


def test_http_docx_to_export_and_config():
    doc = Document()
    for line in SAMPLE.splitlines():
        doc.add_paragraph(line)
    buff = BytesIO()
    doc.save(buff)
    response = client.post("/api/extract", files={"file": ("resume.docx", buff.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")})
    assert response.status_code == 200, response.text
    parsed = response.json()
    assert parsed["resume"]["name"] == "陈晓明"
    assert parsed["source_text"].startswith("陈晓明")
    assert client.get("/api/config").json()["ai_available"] in [True, False]
    enhanced = client.post("/api/enhance", json={"resume": parsed["resume"], "use_ai": False})
    assert enhanced.status_code == 200, enhanced.text
    for kind, mime in [("docx", "wordprocessingml"), ("pdf", "application/pdf")]:
        output = client.post(f"/api/export?format={kind}", json={"resume": enhanced.json()["resume"]})
        assert output.status_code == 200, output.text if output.status_code != 200 else ""
        assert mime in output.headers["content-type"]
        assert len(output.content) > 100


def test_export_blocks_empty_resume():
    response = client.post("/api/export?format=pdf", json={"resume": Resume().model_dump()})
    assert response.status_code == 422


def test_ai_requested_without_credentials_is_not_silently_local(monkeypatch):
    monkeypatch.delenv("PAPERPRO_OPENAI_API_KEY", raising=False)
    response = client.post("/api/enhance", json={"resume": parse_resume(SAMPLE).model_dump(), "use_ai": True})
    assert response.status_code == 400

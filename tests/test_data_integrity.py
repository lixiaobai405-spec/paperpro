from io import BytesIO

from docx import Document
from docx.oxml import OxmlElement
from docx.shared import Inches
from PIL import Image, ImageDraw, ImageFont

from app.parser import extract_file, parse_resume


def test_decimal_numbers_not_stripped_as_numbered_list():
    raw = """李华
其他信息
12.5% 成本改善
1.2亿元 项目预算
3.5年 软件经验
"""
    resume = parse_resume(raw)
    assert "12.5%" in str(resume.additional)
    assert "1.2亿元" in str(resume.additional)
    assert "3.5年" in str(resume.additional)


def test_docx_textbox_content_survives_extraction():
    document = Document()
    document.add_paragraph("Robin Test")
    paragraph = document.add_paragraph()
    container = OxmlElement("w:txbxContent")
    inner_p = OxmlElement("w:p")
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = "TEXTBOX UNIQUE SKILL"
    run.append(text)
    inner_p.append(run)
    container.append(inner_p)
    paragraph._p.append(container)
    stream = BytesIO()
    document.save(stream)
    extracted, _, _ = extract_file("textbox.docx", stream.getvalue())
    assert "Robin Test" in extracted
    assert "TEXTBOX UNIQUE SKILL" in extracted


def test_docx_textbox_break_preserves_separate_numeric_claims():
    doc = Document()
    doc.add_paragraph("Leo")
    paragraph = doc.add_paragraph()
    box = OxmlElement("w:txbxContent")
    content = OxmlElement("w:p")
    run = OxmlElement("w:r")
    first = OxmlElement("w:t")
    first.text = "Team size: 12"
    line_break = OxmlElement("w:br")
    second = OxmlElement("w:t")
    second.text = "5% cost reduction"
    run.extend([first, line_break, second])
    content.append(run)
    box.append(content)
    paragraph._p.append(box)
    stream = BytesIO()
    doc.save(stream)
    extracted, _, _ = extract_file("multiple-lines.docx", stream.getvalue())
    assert "Team size: 12\n5% cost reduction" in extracted
    assert "125% cost reduction" not in extracted


def test_embedded_image_in_word_is_ocr_read_or_warned():
    doc = Document()
    doc.add_paragraph("Terry Smith")
    img = Image.new("RGB", (1300, 260), "white")
    ImageDraw.Draw(img).text((35, 48), "EMBEDDED SKILL", font=ImageFont.truetype("C:/Windows/Fonts/arial.ttf", 68), fill="black")
    data = BytesIO()
    img.save(data, "PNG")
    data.seek(0)
    doc.add_picture(data, width=Inches(5))
    stream = BytesIO()
    doc.save(stream)
    text, warnings, method = extract_file("with-image.docx", stream.getvalue())
    assert "Terry Smith" in text
    assert warnings, "Word image content must never be silently ignored"
    assert "EMBEDDED" in text.upper()
    assert "ocr" in method.lower()

from __future__ import annotations

from io import BytesIO
import os

import pymupdf
import pytest
from docx import Document
from fastapi.testclient import TestClient
from PIL import Image, ImageDraw, ImageFont

from app.enhancer import validate_ai_rewrite
from app.exporter import _fallback_pdf, make_pdf
from app.main import app
from app.models import Entry, Resume
from app.parser import extract_file


def test_pdf_fallback_bmp_chinese_is_searchable(monkeypatch):
    import app.exporter as exporter
    def word_unavailable(_):
        raise RuntimeError("Word temporarily unavailable")
    monkeypatch.setattr(exporter, "_word_pdf", word_unavailable)
    pdf = make_pdf(Resume(name="陈晓明", summary="熟悉Python服务开发"))
    page = pymupdf.open(stream=pdf, filetype="pdf")[0]
    assert "陈晓明" in page.get_text()
    assert abs(page.rect.width - 595.3) < 1
    assert abs(page.rect.height - 841.9) < 1


def test_ai_signed_metric_change_is_flagged_for_review_not_applied():
    from app.enhancer import _changes
    source = Resume(name="Alex", summary="Revenue declined -20%")
    proposal = {"summary": "Revenue declined 20%", "experience_bullets": [], "project_bullets": []}
    suggested = validate_ai_rewrite(source, proposal)
    changes = _changes(source, suggested)
    assert len(changes) == 1
    assert any("数字" in flag for flag in changes[0]["review_flags"])
    assert source.summary == "Revenue declined -20%"


def test_mocked_ai_happy_path_no_fact_added(monkeypatch):
    import app.enhancer as enhancer
    seen = []

    class Response:
        def raise_for_status(self):
            pass
        def json(self):
            return {"choices": [{"message": {"content": (
                '{"summary": "Built a backend service", '
                '"experience_bullets": [["Reduced response latency by 20%"]], '
                '"project_bullets": []}'
            )}}]}

    class Client:
        def __init__(self, *args, **kwargs):
            pass
        def __enter__(self):
            return self
        def __exit__(self, *args):
            pass
        def post(self, path, **kwargs):
            seen.append((path, kwargs))
            return Response()

    monkeypatch.setattr(enhancer.httpx, "Client", Client)
    raw = Resume(name="Alex", email="alex@example.com", summary="Developed backend service",
                 experience=[Entry(heading="Example Corp | 2024", bullets=["Reduced response time 20%"])])
    response = TestClient(app).post("/api/enhance", json={
        "resume": raw.model_dump(), "use_ai": True, "api_key": "TEST-KEY",
        "target_role": "Backend developer"
    })
    assert response.status_code == 200, response.text
    data = response.json()
    assert data["mode"] == "ai"
    assert data["resume"]["name"] == raw.name
    assert data["resume"]["email"] == raw.email
    assert len(data["changes"]) == 2
    assert data["resume"]["experience"][0]["bullets"][0].endswith("20%")
    assert seen[0][0].startswith("https://")
    assert seen[0][1]["headers"]["Authorization"] == "Bearer TEST-KEY"
    assert any("Backend developer" in x["content"] for x in seen[0][1]["json"]["messages"])


def test_uploaded_encrypted_pdf_gives_controlled_error():
    document = pymupdf.open()
    document.new_page().insert_text((30, 30), "Private resume")
    encrypted = document.tobytes(encryption=pymupdf.PDF_ENCRYPT_AES_256, owner_pw="owner", user_pw="password")
    response = TestClient(app).post("/api/extract", files={
        "file": ("locked.pdf", encrypted, "application/pdf")
    })
    assert response.status_code == 422
    assert "PDF" in response.json()["detail"]


def test_upload_over_page_limit():
    document = pymupdf.open()
    for _ in range(26):
        document.new_page().insert_text((10, 30), "More")
    response = TestClient(app).post("/api/extract", files={
        "file": ("huge.pdf", document.tobytes(), "application/pdf")
    })
    assert response.status_code == 422
    assert "25" in response.json()["detail"]


def test_file_mime_mismatch_rejected():
    response = TestClient(app).post("/api/extract", files={
        "file": ("photo.png", b"Not a png", "image/png")
    })
    assert response.status_code == 422


@pytest.mark.skipif(os.name != "nt", reason="Legacy Word COM conversion is Windows-only")
def test_legacy_doc_end_to_end(tmp_path):
    import pythoncom
    import win32com.client

    source = tmp_path / "legacy-source.docx"
    legacy = tmp_path / "legacy-export.doc"
    document = Document()
    document.add_paragraph("LEGACY CANDIDATE")
    document.add_paragraph("Some Legacy Experience")
    document.save(source)

    pythoncom.CoInitialize()
    word = None
    opened = None
    try:
        word = win32com.client.DispatchEx("Word.Application")
        word.Visible = False
        word.DisplayAlerts = 0
        word.AutomationSecurity = 3
        opened = word.Documents.Open(str(source), ReadOnly=False, AddToRecentFiles=False)
        opened.SaveAs(str(legacy), FileFormat=0)
        opened.Close(SaveChanges=False)
        opened = None
    finally:
        if opened is not None:
            opened.Close(SaveChanges=False)
        if word is not None:
            word.Quit()
        pythoncom.CoUninitialize()

    raw, warnings, method = extract_file("old.doc", legacy.read_bytes())
    assert "LEGACY CANDIDATE" in raw
    assert method == "word-doc"

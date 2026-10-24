from io import BytesIO
import os
from types import SimpleNamespace
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
from PIL import Image
import pymupdf
import pytest

from app import exporter, parser
from app.document_converter import ConversionError, ConverterUnavailable
from app.models import Resume


def linux_module(monkeypatch, module):
    # Replace this module's os binding, not the process-wide os.name (which
    # would make pathlib choose an incompatible concrete Path class).
    monkeypatch.setattr(module, "os", SimpleNamespace(name="posix", environ=os.environ))


def legacy_input():
    return bytes.fromhex("D0CF11E0A1B11AE1") + b"synthetic DOC"


def test_linux_doc_reuses_docx_table_and_image_extraction(monkeypatch):
    linux_module(monkeypatch, parser)
    doc = Document()
    doc.add_paragraph("陈晓明")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "工作经历"
    table.cell(0, 1).text = "示范科技 Python 工程师"
    image = BytesIO()
    Image.new("RGB", (100, 60), "white").save(image, "PNG")
    image.seek(0)
    doc.add_picture(image)
    blob = BytesIO()
    doc.save(blob)
    monkeypatch.setattr(parser, "convert_document", lambda data, source, target: blob.getvalue())
    monkeypatch.setattr(parser, "_ocr_image", lambda image: "IMAGE EVIDENCE")
    text, warnings, method = parser.extract_file("legacy.doc", legacy_input())
    assert "陈晓明" in text and "示范科技 Python 工程师" in text
    assert "IMAGE EVIDENCE" in text
    assert method == "libreoffice-doc+ocr"
    assert any("LibreOffice" in warning for warning in warnings)
    assert any("OCR" in warning for warning in warnings)


def test_linux_doc_rejects_conversion_expansion_before_parsing(monkeypatch):
    linux_module(monkeypatch, parser)
    blob = BytesIO()
    with ZipFile(blob, "w", ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", "x")
        archive.writestr("word/document.xml", b"x" * (48 * 1024 * 1024 + 1))
    monkeypatch.setattr(parser, "convert_document", lambda *args: blob.getvalue())
    with pytest.raises(ConversionError, match="安全限制"):
        parser.extract_file("large.doc", legacy_input())


def test_linux_doc_missing_office_is_not_empty_success(monkeypatch):
    linux_module(monkeypatch, parser)

    def unavailable(*args):
        raise ConverterUnavailable("请安装 LibreOffice")

    monkeypatch.setattr(parser, "convert_document", unavailable)
    with pytest.raises(ConverterUnavailable, match="LibreOffice"):
        parser.extract_file("old.doc", legacy_input())


def test_linux_pdf_uses_same_docx_template(monkeypatch):
    linux_module(monkeypatch, exporter)
    monkeypatch.setenv("PAPERPRO_PDF_ENGINE", "auto")
    seen = []
    with pymupdf.open() as pdf:
        pdf.new_page().insert_text((40, 50), "Converted synthetic resume")
        expected = pdf.tobytes()

    def convert(data, source, target):
        seen.append((Document(BytesIO(data)).paragraphs[0].text, source, target))
        return expected

    monkeypatch.setattr(exporter, "convert_document", convert)
    output = exporter.make_pdf(Resume(name="陈晓明", summary="Python 后端开发"))
    assert output == expected
    assert seen == [("陈晓明", "docx", "pdf")]


def test_strict_libreoffice_failure_does_not_silently_change_layout(monkeypatch):
    monkeypatch.setenv("PAPERPRO_PDF_ENGINE", "libreoffice")

    def fail(*args):
        raise ConversionError("文档转换超时")

    monkeypatch.setattr(exporter, "convert_document", fail)
    monkeypatch.setattr(exporter, "_fallback_pdf", lambda _: pytest.fail("strict mode must not fall back"))
    with pytest.raises(ConversionError, match="超时"):
        exporter.make_pdf(Resume(name="陈晓明", summary="测试"))


def test_auto_linux_can_use_explicitly_bounded_fallback(monkeypatch):
    linux_module(monkeypatch, exporter)
    monkeypatch.setenv("PAPERPRO_PDF_ENGINE", "auto")

    def unavailable(*args):
        raise ConverterUnavailable("LibreOffice unavailable")

    monkeypatch.setattr(exporter, "convert_document", unavailable)
    output = exporter.make_pdf(Resume(name="陈晓明", summary="本地备用导出"))
    with pymupdf.open(stream=output, filetype="pdf") as pdf:
        assert "陈晓明" in pdf[0].get_text()


def test_fallback_keeps_unsupported_character_guard(monkeypatch):
    monkeypatch.setenv("PAPERPRO_PDF_ENGINE", "reportlab")
    with pytest.raises(ValueError, match="罕见汉字"):
        exporter.make_pdf(Resume(name="𠮷", summary="不能静默丢字"))


def test_linux_template_font_and_override(monkeypatch):
    linux_module(monkeypatch, exporter)
    monkeypatch.delenv("PAPERPRO_RESUME_FONT", raising=False)
    doc = Document(BytesIO(exporter.make_docx(Resume(name="中文字体"))))
    assert doc.paragraphs[0].runs[0].font.name == "Noto Sans CJK SC"
    assert doc.styles["Normal"].font.name == "Noto Sans CJK SC"
    monkeypatch.setenv("PAPERPRO_RESUME_FONT", "WenQuanYi Zen Hei")
    doc = Document(BytesIO(exporter.make_docx(Resume(name="中文字体"))))
    assert doc.paragraphs[0].runs[0].font.name == "WenQuanYi Zen Hei"


@pytest.mark.parametrize("engine", ["invalid", "WORD "])
def test_unknown_pdf_engine_fails_clearly(monkeypatch, engine):
    monkeypatch.setenv("PAPERPRO_PDF_ENGINE", engine)
    with pytest.raises(ValueError, match="PAPERPRO_PDF_ENGINE"):
        exporter.make_pdf(Resume(name="Test", summary="Synthetic"))

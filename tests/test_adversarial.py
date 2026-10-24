from io import BytesIO

import pymupdf
import pytest
from docx import Document
from docx.shared import Mm
from fastapi.testclient import TestClient
from PIL import Image, ImageDraw, ImageFont

from app.main import app
from app.models import Entry, Resume
from app.parser import extract_file
from app.enhancer import ai_enhance, local_enhance, validate_ai_rewrite
from app.exporter import make_docx, make_pdf, _fallback_pdf


def test_mixed_native_and_scanned_pdf_keeps_both():
    img = Image.new("RGB", (1300, 320), "white")
    ImageDraw.Draw(img).text((30, 55), "LUMINOUS COMPANY", font=ImageFont.truetype("C:/Windows/Fonts/arial.ttf", 68), fill="black")
    buf = BytesIO()
    img.save(buf, format="PNG")
    pdf = pymupdf.open()
    page = pdf.new_page(width=750, height=800)
    page.insert_text((40, 60), "This is a long native page heading with text", fontsize=17)
    page.insert_image(pymupdf.Rect(40, 140, 710, 360), stream=buf.getvalue())
    text, warnings, method = extract_file("mixed.pdf", pdf.tobytes())
    assert "heading" in text.lower()
    assert "LUMINOUS" in text.upper()
    assert "ocr" in method.lower()
    assert warnings


def test_docx_headers_nested_tables_and_footers_are_not_discarded():
    document = Document()
    document.sections[0].header.paragraphs[0].text = "unique-header-contact@demo.com"
    document.add_paragraph("Chen Mei")
    outer = document.add_table(rows=1, cols=1)
    inner = outer.cell(0, 0).add_table(rows=1, cols=1)
    inner.cell(0, 0).text = "UNIQUE NESTED ORGANIZATION"
    document.sections[0].footer.paragraphs[0].text = "CERTIFICATION-ISO"
    blob = BytesIO()
    document.save(blob)
    text, _, _ = extract_file("candidate.docx", blob.getvalue())
    assert "unique-header-contact@demo.com" in text
    assert "UNIQUE NESTED ORGANIZATION" in text
    assert "CERTIFICATION-ISO" in text


def test_cleanup_never_mistakes_negative_values_for_bullet_prefix():
    source = Resume(
        name="Test", summary="-20% revenue", experience=[Entry(heading="X", bullets=["-15% margin"])],
        projects=[Entry(heading="P", bullets=["-12 failures"])]
    )
    cleaned, _ = local_enhance(source)
    assert cleaned.summary == source.summary
    assert cleaned.experience[0].bullets[0] == source.experience[0].bullets[0]
    assert cleaned.projects[0].bullets[0] == source.projects[0].bullets[0]


def test_changed_numbers_are_reviewable_model_suggestions_not_global_422():
    from app.enhancer import _changes
    source = Resume(name="Test", summary="Reduced 20% errors and 90% delays")
    rewrite = {"summary": "Reduced 90% errors and 20% delays", "experience_bullets": [], "project_bullets": []}
    result = validate_ai_rewrite(source, rewrite)
    assert result.summary == rewrite["summary"]
    change = _changes(source, result)[0]
    assert any("数字" in flag for flag in change["review_flags"])
    assert source.summary == "Reduced 20% errors and 90% delays"
    rewrite["summary"] = "Reduced 20% errors and 90% delays, with 20% cost reduction"
    change = _changes(source, validate_ai_rewrite(source, rewrite))[0]
    assert change["review_flags"]


def test_ai_new_technical_claim_is_flagged_for_explicit_human_review():
    source = Resume(
        name="虚构候选人",
        summary="熟悉 Python 开发",
        experience=[Entry(heading="已确认单位｜后端开发", bullets=["使用 Python 开发 API"])]
    )
    proposal = {
        "summary": source.summary,
        "experience_bullets": [["使用 Python 与 Kubernetes 开发 API"]],
        "project_bullets": []
    }
    # The model may suggest this wording, but must not silently apply it.
    from app.enhancer import _changes
    result = validate_ai_rewrite(source, proposal)
    changes = _changes(source, result)
    assert len(changes) == 1
    assert any("技术" in warning for warning in changes[0]["review_flags"])
    assert source.experience[0].bullets[0] == "使用 Python 开发 API"


def test_ai_can_rephrase_existing_technical_skill_without_new_claim():
    source = Resume(
        name="虚构候选人",
        summary="熟悉 Python 开发",
        experience=[Entry(heading="已确认单位｜后端开发", bullets=["使用 Python 开发 API"])]
    )
    proposal = {
        "summary": "具有 Python 开发经验",
        "experience_bullets": [["使用 Python 实现 API 接口"]],
        "project_bullets": []
    }
    output = validate_ai_rewrite(source, proposal)
    assert output.experience[0].bullets[0] == "使用 Python 实现 API 接口"


def test_loose_rewrite_shape_skips_unmatched_entries_but_keeps_other_changes():
    original = Resume(name="测试", summary="松散描述", experience=[
        Entry(heading="临时项目", bullets=["负责测试", "整理文档"]),
        Entry(heading="其他经历", bullets=["支持发布"]),
    ])
    proposal = {
        "summary": "描述项目测试与文档整理",
        "experience_bullets": [
            ["测试", "整理", "新增了一条"],  # Cannot align: preserve original.
            ["参与发布支持"],
        ],
        "project_bullets": [],
        "name": "模型编造的人名",  # Outside writable fields.
    }
    result = validate_ai_rewrite(original, proposal)
    assert result.name == "测试"
    assert result.summary == "描述项目测试与文档整理"
    assert result.experience[0].bullets == ["负责测试", "整理文档"]
    assert result.experience[1].bullets == ["参与发布支持"]
    assert original.experience[1].bullets == ["支持发布"]


def test_docx_page_size_is_a4():
    output = make_docx(Resume(name="Name", summary="Work"))
    document = Document(BytesIO(output))
    for section in document.sections:
        assert abs(section.page_width.mm - 210) < 1
        assert abs(section.page_height.mm - 297) < 1


def test_plaintext_key_never_sent_to_external_host_with_localhost_userinfo(monkeypatch):
    monkeypatch.setenv("PAPERPRO_OPENAI_BASE_URL", "http://localhost:80@remote.example/v1")
    with pytest.raises(ValueError):
        ai_enhance(Resume(name="X", summary="Hello"), api_key="a-sensitive-key")


def test_request_validation_error_does_not_echo_api_key():
    response = TestClient(app).post("/api/enhance", json={"api_key": "SENSITIVE-KEY-ABC123", "use_ai": True})
    assert response.status_code == 422
    assert "SENSITIVE-KEY-ABC123" not in response.text


def test_pdf_word_engine_preserves_extended_han_character():
    import pymupdf
    pdf = make_pdf(Resume(name="𠮷野", summary="Software developer"))
    text = "\n".join(page.get_text() for page in pymupdf.open(stream=pdf, filetype="pdf"))
    assert "𠮷野" in text


def test_pdf_fallback_rejects_rare_glyph_instead_of_corrupting_name():
    with pytest.raises(ValueError, match="罕见汉字"):
        _fallback_pdf(Resume(name="𠮷野", summary="Software developer"))


def test_ai_null_content_is_safe_error(monkeypatch):
    class FakeResponse:
        def raise_for_status(self):
            pass
        def json(self):
            return {"choices": [{"message": {"content": None}}]}
    class FakeClient:
        def __init__(self, *args, **kwargs):
            pass
        def __enter__(self):
            return self
        def __exit__(self, *args):
            pass
        def post(self, *args, **kwargs):
            return FakeResponse()
    monkeypatch.setattr("app.enhancer.httpx.Client", FakeClient)
    with pytest.raises(RuntimeError):
        ai_enhance(Resume(name="Demo", summary="Built APIs"), api_key="fake")


def test_chinese_scanned_image_ocr():
    img = Image.new("RGB", (1600, 330), "white")
    font = ImageFont.truetype("C:/Windows/Fonts/msyh.ttc", 82)
    ImageDraw.Draw(img).text((40, 60), "软件开发工程师", font=font, fill="black")
    buf = BytesIO()
    img.save(buf, format="PNG")
    content, warnings, method = extract_file("chinese.png", buf.getvalue())
    assert "软件开发" in content
    assert method == "image-ocr"

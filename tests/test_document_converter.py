from io import BytesIO
from pathlib import Path
import subprocess
import threading
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
import pymupdf
import pytest

from app import document_converter as converter


def docx_bytes():
    stream = BytesIO()
    doc = Document()
    doc.add_paragraph("陈晓明 synthetic resume")
    doc.save(stream)
    return stream.getvalue()


def pdf_bytes():
    with pymupdf.open() as pdf:
        pdf.new_page().insert_text((50, 60), "Synthetic resume")
        return pdf.tobytes()


@pytest.fixture
def office(monkeypatch):
    monkeypatch.setattr(converter, "find_libreoffice", lambda: "/usr/bin/libreoffice")


def install_output(monkeypatch, output):
    seen = []

    def run(command, environment):
        seen.append((command, environment, Path(command[-1]).parent))
        extension = command[command.index("--convert-to") + 1].split(":")[0]
        destination = Path(command[command.index("--outdir") + 1]) / ("input." + extension)
        destination.write_bytes(output)

    monkeypatch.setattr(converter, "_run_conversion", run)
    return seen


def test_conversion_uses_private_profile_and_does_not_inherit_secrets(office, monkeypatch):
    monkeypatch.setenv("PAPERPRO_DEEPSEEK_API_KEY", "private-test-key")
    monkeypatch.setenv("AWS_SECRET_ACCESS_KEY", "private-test-value")
    profiles = []

    def run(command, environment):
        root = Path(command[-1]).parent
        profile = root / "profile"
        configuration = (profile / "user" / "registrymodifications.xcu").read_text()
        assert "DisableMacrosExecution" in configuration
        assert "DisableActiveContent" in configuration
        assert "BlockUntrustedRefererLinks" in configuration
        assert '<value>2</value>' in configuration  # Never update external links.
        assert not any("KEY" in name or "TOKEN" in name for name in environment)
        assert "private-test" not in str(environment)
        assert command[-1] == str(root / "input.doc")
        assert any(arg.startswith("-env:UserInstallation=file:") for arg in command)
        assert "--headless" in command
        profiles.append(root)
        (root / "output" / "input.docx").write_bytes(docx_bytes())

    monkeypatch.setattr(converter, "_run_conversion", run)
    first = converter.convert_document(b"synthetic DOC", "doc", "docx")
    second = converter.convert_document(b"synthetic DOC", "doc", "docx")
    assert "陈晓明" in Document(BytesIO(first)).paragraphs[0].text
    assert second.startswith(b"PK")
    assert profiles[0] != profiles[1]
    assert all(not folder.exists() for folder in profiles)


def test_pdf_conversion_validates_readable_document(office, monkeypatch):
    expected = pdf_bytes()
    seen = install_output(monkeypatch, expected)
    assert converter.convert_document(docx_bytes(), "docx", "pdf") == expected
    assert not seen[0][2].exists()


@pytest.mark.parametrize("content", [b"", b"%PDF-not-a-real-document", b"plain text"])
def test_bad_pdf_output_is_not_success(office, monkeypatch, content):
    seen = install_output(monkeypatch, content)
    with pytest.raises(converter.ConversionError):
        converter.convert_document(docx_bytes(), "docx", "pdf")
    assert not seen[0][2].exists()


def test_missing_output_is_controlled_error_and_releases_slot(office, monkeypatch):
    slots = threading.BoundedSemaphore(1)
    monkeypatch.setattr(converter, "_conversion_slots", slots)
    roots = []
    monkeypatch.setattr(converter, "_run_conversion", lambda command, env: roots.append(Path(command[-1]).parent))
    with pytest.raises(converter.ConversionError, match="生成"):
        converter.convert_document(docx_bytes(), "docx", "pdf")
    assert slots.acquire(blocking=False)
    slots.release()
    assert not roots[0].exists()


def test_missing_office_is_actionable(monkeypatch):
    monkeypatch.setattr(converter, "find_libreoffice", lambda: None)
    with pytest.raises(converter.ConverterUnavailable, match="LibreOffice"):
        converter.convert_document(b"document", "doc", "docx")


def test_saturated_converter_returns_bounded_error(office, monkeypatch):
    slots = threading.BoundedSemaphore(1)
    slots.acquire()
    monkeypatch.setattr(converter, "_conversion_slots", slots)
    monkeypatch.setattr(converter, "QUEUE_TIMEOUT_SECONDS", 0.01)
    with pytest.raises(converter.ConversionError, match="繁忙"):
        converter.convert_document(b"document", "doc", "docx")
    slots.release()


def test_source_and_output_size_limits(office, monkeypatch):
    monkeypatch.setattr(converter, "MAX_SOURCE_BYTES", 16)
    with pytest.raises(converter.ConversionError, match="大小"):
        converter.convert_document(b"x" * 17, "doc", "docx")
    monkeypatch.setattr(converter, "MAX_OUTPUT_BYTES", 16)
    install_output(monkeypatch, b"x" * 17)
    with pytest.raises(converter.ConversionError, match="大小"):
        converter.convert_document(b"document", "doc", "docx")


@pytest.mark.parametrize("source,target", [("../../doc", "pdf"), ("docx", "sh"), ("pdf", "docx")])
def test_formats_cannot_inject_paths_or_arguments(source, target):
    with pytest.raises(converter.ConversionError, match="格式"):
        converter.convert_document(b"document", source, target)


def test_docx_limits_apply_to_converted_archives():
    blob = BytesIO()
    with ZipFile(blob, "w", ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", "x")
        archive.writestr("word/document.xml", b"x" * (48 * 1024 * 1024 + 1))
    with pytest.raises(converter.ConversionError, match="安全限制"):
        converter.validate_docx(blob.getvalue())


def test_docx_requires_word_document():
    blob = BytesIO()
    with ZipFile(blob, "w") as archive:
        archive.writestr("arbitrary.txt", "not a Word document")
    with pytest.raises(converter.ConversionError, match="DOCX"):
        converter.validate_docx(blob.getvalue())


def test_subprocess_timeout_stops_process_and_hides_document_output(monkeypatch):
    seen = {}

    class HangingProcess:
        pid = 98765

        def wait(self, timeout=None):
            raise subprocess.TimeoutExpired("office", timeout)

    def popen(command, **kwargs):
        seen.update(kwargs)
        assert isinstance(command, list)
        return HangingProcess()

    killed = []
    monkeypatch.setattr(converter.subprocess, "Popen", popen)
    monkeypatch.setattr(converter, "_terminate_process", lambda process: killed.append(process.pid))
    with pytest.raises(converter.ConversionError, match="超时"):
        converter._run_conversion(["office", "input.docx"], {"PATH": "test"})
    assert killed == [98765]
    assert seen["stdout"] == subprocess.DEVNULL
    assert seen["stderr"] == subprocess.DEVNULL
    assert seen["stdin"] == subprocess.DEVNULL
    assert seen["shell"] is False


def test_nonzero_exit_is_not_reported_as_success(monkeypatch):
    class FailedProcess:
        def wait(self, timeout=None):
            return 1

    monkeypatch.setattr(converter.subprocess, "Popen", lambda *args, **kwargs: FailedProcess())
    with pytest.raises(converter.ConversionError, match="失败"):
        converter._run_conversion(["office", "input.docx"], {})

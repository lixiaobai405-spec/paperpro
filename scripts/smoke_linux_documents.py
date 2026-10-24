"""Offline Linux image acceptance using synthetic documents, never an AI API."""
from __future__ import annotations

import argparse
from concurrent.futures import ThreadPoolExecutor
from io import BytesIO
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
import json
import os
from pathlib import Path
import re
import subprocess
import sys
import tempfile
import threading
import time
from xml.etree import ElementTree

from docx import Document
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.shared import Inches
from fastapi.testclient import TestClient
from PIL import Image, ImageDraw, ImageFont
import pymupdf

from app import document_converter as office
from app.exporter import _fallback_pdf, make_docx, make_pdf
from app.main import app
from app.models import Education, Entry, Resume
from app.parser import extract_file
from app.pdf_runtime import serialized_pdf


def sample() -> Resume:
    return Resume(
        name="林知远", title="全栈开发工程师", email="lin.zhiyuan@example.com",
        phone="13800000000", location="杭州",
        summary="虚构测试简历。具有2年Web开发经验，熟悉Python和React，关注测试与交付。",
        skills=["Python", "FastAPI", "React", "SQL", "Docker"],
        experience=[Entry(heading="2024.07 - 2026.06 示例科技 | 全栈开发工程师", bullets=[
            "负责工单接口与前端页面，在固定测试集上将响应时间从820ms降至310ms。",
            "编写45项接口测试及12条浏览器回归用例，保留原始数字与事实。",
        ])],
        projects=[Entry(heading="校园资料检索助手", bullets=["实现文件上传、关键词检索和原文核对。 "])],
        education=[Education(heading="2020.09 - 2024.06 示例大学 | 计算机科学与技术 本科")],
        additional=["合成测试数据，不代表真实候选人的学历、经历或成果。"],
    )


@serialized_pdf
def normalized_pdf(data: bytes) -> str:
    with pymupdf.open(stream=data, filetype="pdf") as document:
        return re.sub(r"\s", "", "".join(page.get_text() for page in document))


def legacy_fixture(docx: bytes) -> bytes:
    with tempfile.TemporaryDirectory(prefix="paperpro-doc-fixture-") as folder:
        root = Path(folder)
        source = root / "synthetic.docx"
        source.write_bytes(docx)
        profile = root / "profile"
        (profile / "user").mkdir(parents=True)
        (profile / "user" / "registrymodifications.xcu").write_text(office._PROFILE, encoding="utf-8")
        office._run_conversion([
            office.find_libreoffice(), "-env:UserInstallation=" + profile.as_uri(),
            "--headless", "--norestore", "--convert-to", "doc:MS Word 97",
            "--outdir", str(root), str(source),
        ], {"PATH": os.environ["PATH"], "HOME": str(root), "LANG": "C.UTF-8"})
        # Read back the real profile after LO has opened/saved it. Combined
        # with the image's installed-schema gate, this verifies that the
        # supported policy lives in the actual user profile, not a sibling.
        configuration = ElementTree.parse(profile / "user" / "registrymodifications.xcu")
        names = {element.get("{http://openoffice.org/2001/registry}name"): element.findtext("value")
                 for element in configuration.iter("prop")}
        assert names["DisableMacrosExecution"] == "true"
        assert names["DisableActiveContent"] == "true"
        assert names["BlockUntrustedRefererLinks"] == "true"
        assert names["Link"] == "2"
        result = (root / "synthetic.doc").read_bytes()
        assert result.startswith(bytes.fromhex("D0CF11E0A1B11AE1"))
        return result


def verify_external_image_is_not_fetched() -> None:
    requests = []
    image_data = BytesIO()
    Image.new("RGB", (30, 30), "white").save(image_data, "PNG")

    class Probe(BaseHTTPRequestHandler):
        def do_GET(self):
            requests.append(self.path)
            self.send_response(200)
            self.send_header("Content-Type", "image/png")
            self.end_headers()
            self.wfile.write(image_data.getvalue())

        def log_message(self, *args):
            pass

    server = ThreadingHTTPServer(("127.0.0.1", 0), Probe)
    thread = threading.Thread(target=server.serve_forever, kwargs={"poll_interval": 0.05}, daemon=True)
    thread.start()
    try:
        document = Document()
        document.add_paragraph("外链阻断测试，使用合成图片")
        image_data.seek(0)
        document.add_picture(image_data)
        relation = document.part.relate_to(
            f"http://127.0.0.1:{server.server_port}/external.png", RELATIONSHIP_TYPE.IMAGE, is_external=True,
        )
        blip = document.element.xpath(".//a:blip")[0]
        del blip.attrib[qn("r:embed")]
        blip.set(qn("r:link"), relation)
        content = BytesIO()
        document.save(content)
        office.convert_document(content.getvalue(), "docx", "pdf")
        assert not requests, "The converter fetched an external image"
    finally:
        server.shutdown()
        server.server_close()
        thread.join(timeout=1)


def verify_timeout_cleanup() -> None:
    with tempfile.TemporaryDirectory(prefix="paperpro-timeout-check-") as folder:
        pid_file = Path(folder) / "child.pid"
        code = (
            "import pathlib,subprocess,sys,time; "
            "p=subprocess.Popen([sys.executable,'-c','import time; time.sleep(60)']); "
            "pathlib.Path(sys.argv[1]).write_text(str(p.pid)); time.sleep(60)"
        )
        previous = office.CONVERSION_TIMEOUT_SECONDS
        office.CONVERSION_TIMEOUT_SECONDS = 1
        try:
            try:
                office._run_conversion([sys.executable, "-c", code, str(pid_file)], {"PATH": os.environ["PATH"]})
            except office.ConversionError as error:
                assert "超时" in str(error)
            else:
                raise AssertionError("A hanging process must time out")
        finally:
            office.CONVERSION_TIMEOUT_SECONDS = previous
        child = int(pid_file.read_text())
        state_file = Path(f"/proc/{child}/stat")
        for _ in range(20):
            if not state_file.exists() or state_file.read_text().split()[2] == "Z":
                break
            time.sleep(0.05)
        else:
            raise AssertionError("Timed-out converter child is still running")


def run(output: Path) -> dict:
    assert os.name == "posix", "Run this acceptance script inside the Linux image"
    assert office.find_libreoffice(), "LibreOffice must be installed"
    assert os.environ.get("PAPERPRO_PDF_ENGINE") == "libreoffice", "Exercise strict production conversion"
    output.mkdir(parents=True, exist_ok=True)
    checks = []
    resume = sample()
    docx = make_docx(resume)
    pdf = make_pdf(resume, docx)
    text = normalized_pdf(pdf)
    for expected in ("林知远", "示例科技", "820ms", "310ms", "示例大学", "45", "12"):
        assert expected in text, f"Missing PDF evidence: {expected}"
    with pymupdf.open(stream=pdf, filetype="pdf") as document:
        assert len(document) == 1
        assert abs(document[0].rect.width - 595.3) < 1
        document[0].get_pixmap(matrix=pymupdf.Matrix(1.6, 1.6)).save(output / "resume-preview.png")
    (output / "resume.docx").write_bytes(docx)
    (output / "resume.pdf").write_bytes(pdf)
    checks.append("Chinese one-page DOCX to PDF with retained facts")

    legacy = legacy_fixture(docx)
    source, warnings, method = extract_file("synthetic.doc", legacy)
    assert "林知远" in source and "示例科技" in source
    assert method == "libreoffice-doc" and warnings
    (output / "legacy-resume.doc").write_bytes(legacy)
    checks.append("Real legacy DOC to DOCX to existing parser")
    checks.append("Installed-policy profile read back after real LibreOffice conversion")
    verify_external_image_is_not_fetched()
    checks.append("External DOCX image URL is not fetched (local HTTP probe)")

    # The conversion path must preserve tables and continue to OCR images.
    rich = Document()
    rich.add_paragraph("王小林")
    row = rich.add_table(rows=1, cols=2).rows[0]
    row.cells[0].text = "工作经历"
    row.cells[1].text = "示例企业 FastAPI 工程师"
    image = Image.new("RGB", (1300, 220), "white")
    font = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", 65)
    ImageDraw.Draw(image).text((30, 55), "EMBEDDED SKILL", fill="black", font=font)
    picture = BytesIO()
    image.save(picture, "PNG")
    picture.seek(0)
    rich.add_picture(picture, width=Inches(5))
    rich_data = BytesIO()
    rich.save(rich_data)
    rich_doc = legacy_fixture(rich_data.getvalue())
    rich_text, rich_warnings, rich_method = extract_file("table-image.doc", rich_doc)
    assert "示例企业" in rich_text and "EMBEDDED" in rich_text.upper()
    assert rich_method == "libreoffice-doc+ocr" and rich_warnings
    checks.append("Legacy DOC table and embedded-image OCR retained")

    long_resume = Resume(name="长文档候选人", summary="多页测试", experience=[Entry(
        heading="2024 - 2026 示例公司 | 工程师",
        bullets=[f"第{i:03d}条履历：设计业务流程、编写回归测试并记录验证结果。" for i in range(1, 65)],
    )])
    long_pdf = make_pdf(long_resume)
    long_text = normalized_pdf(long_pdf)
    assert "第001条" in long_text and "第064条" in long_text
    with pymupdf.open(stream=long_pdf, filetype="pdf") as document:
        assert len(document) >= 2
        for index, page in enumerate(document):
            assert abs(page.rect.width - 595.3) < 1
            for block in page.get_text("blocks"):
                assert block[0] >= 0 and block[1] >= 0
                assert block[2] <= page.rect.width + 1 and block[3] <= page.rect.height + 1
            if index in {0, len(document) - 1}:
                page.get_pixmap(matrix=pymupdf.Matrix(1.3, 1.3)).save(output / f"multipage-{index + 1}.png")
    (output / "multipage.pdf").write_bytes(long_pdf)
    checks.append("A4 multipage export retains first and last numbered facts")

    rare_pdf = make_pdf(Resume(name="𠮷田", summary="罕见汉字专项测试"))
    assert "𠮷田" in normalized_pdf(rare_pdf), "Supplementary CJK must not silently disappear"
    (output / "supplementary-cjk.pdf").write_bytes(rare_pdf)
    with pymupdf.open(stream=rare_pdf, filetype="pdf") as document:
        document[0].get_pixmap(matrix=pymupdf.Matrix(1.6, 1.6)).save(output / "supplementary-cjk-preview.png")
    checks.append("Supplementary CJK searchable in LibreOffice PDF")

    def concurrent_job(number):
        data = make_pdf(Resume(name=f"并发候选人{number}", summary="独立临时文件测试"))
        assert f"并发候选人{number}" in normalized_pdf(data)

    with ThreadPoolExecutor(max_workers=3) as pool:
        list(pool.map(concurrent_job, range(3)))
    checks.append("Three concurrent conversions retain separate candidates")
    verify_timeout_cleanup()
    checks.append("Real process-group timeout terminates converter children")

    fallback = _fallback_pdf(resume)
    assert "林知远" in normalized_pdf(fallback)
    with pymupdf.open(stream=fallback, filetype="pdf") as document:
        assert any(item[1] in {"ttf", "otf"} for item in document.get_page_fonts(0))
    checks.append("Chinese fallback uses an embedded font")

    with TestClient(app) as client:
        assert client.get("/api/health").status_code == 200
        assert "<html" in client.get("/").text.lower()
        assert client.get("/classic").status_code == 200
        uploaded = client.post("/api/extract", files={"file": ("synthetic.doc", legacy, "application/msword")})
        assert uploaded.status_code == 200, uploaded.text
        assert uploaded.json()["resume"]["name"] == "林知远"
        diagnosis = client.post("/api/diagnosis/parse", files={"file": ("synthetic.doc", legacy, "application/msword")})
        assert diagnosis.status_code == 200, diagnosis.text
        response = client.post("/api/export?format=pdf", json={"resume": resume.model_dump()})
        assert response.status_code == 200 and "林知远" in normalized_pdf(response.content)
        response = client.post("/api/export?format=docx", json={"resume": resume.model_dump()})
        assert response.status_code == 200 and response.content.startswith(b"PK")
    checks.append("Both DOC upload APIs, built frontend and DOCX/PDF HTTP export")

    report = {
        "status": "passed", "checks": checks, "count": len(checks),
        "libreoffice": subprocess.check_output([office.find_libreoffice(), "--version"], text=True).strip(),
        "uid": os.getuid(), "ai_calls": 0,
        "limits": "Synthetic fixtures only; not real-resume accuracy or Word pixel equivalence",
    }
    (output / "report.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    return report


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, default=Path("/tmp/paperpro-linux-smoke"))
    arguments = parser.parse_args()
    print(json.dumps(run(arguments.output), ensure_ascii=False, indent=2))
